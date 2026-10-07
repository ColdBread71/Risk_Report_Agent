"""Presentation-layer exports for the current Node [0]-[3] TARA state."""

import json
import math
import re
from datetime import UTC, datetime
from enum import Enum
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.logger import setup_logger
from core.workflow_status import ASSET_REVIEWABLE_STATUSES, is_asset_reviewable_state
from schemas.evidence import EvidenceRef
from schemas.product import AssetIdentificationResult


logger = setup_logger("exporter")

ASSET_REVIEW_SCHEMA_VERSION = "node_0_2_review.v2"
PROGRESS_SCHEMA_VERSION = "workflow_progress.v2"

_ASSET_REVIEW_SHEETS = {
    "Node0 范围定界",
    "Node1 产品概况",
    "Node1 产品功能",
    "Node1 组件清单",
    "Node1 通信矩阵",
    "Node2 资产与目标",
    "方法与准则",
    "证据追溯",
}

_ASSET_HEADERS = (
    "资产编号",
    "资产名称",
    "资产类型",
    "资产状态",
    "关联功能编号",
    "关联功能名称",
    "功能关系与依据",
    "所在位置",
    "资产价值或受损后果",
    "资产说明",
    "目标编号",
    "目标类别",
    "目标依据类型",
    "目标支持状态",
    "网络安全目标",
    "目标法律与标准依据",
)


def _valid_artifact_stem(artifact_stem: str, *, asset_review: bool = False) -> bool:
    if artifact_stem == "report":
        return True
    node_pattern = r"node2(?:_[1-9]\d*)?" if asset_review else r"node[012](?:_[1-9]\d*)?"
    return re.fullmatch(node_pattern, artifact_stem) is not None

_DISPLAY_TRANSLATIONS = {
    "PASS": "通过",
    "PARTIAL": "部分满足",
    "FAIL": "不通过",
    "Default": "默认类",
    "Important": "重要类",
    "Critical": "关键类",
    "confidentiality": "保密性",
    "integrity": "完整性",
    "availability": "可用性",
    "authenticity": "真实性",
    "accountability": "可追责性",
    "data_minimization": "数据最小化",
    "credential": "凭据与密钥",
    "function": "产品功能",
    "hardware": "硬件",
    "network": "网络",
    "software": "软件与配置",
    "service": "外部服务",
    "data": "数据",
    "user": "用户",
    "property": "财产与生产对象",
    "environment": "环境",
    "public_interest": "公共利益",
    "other": "其他",
    "confirmed": "已确认",
    "conditional": "条件项",
    "external": "产品边界外",
    "unknown": "待确认",
    "external_affected": "外部/受影响对象",
    "needs_confirmation": "待工程师确认",
    "supported": "依据已支持",
    "explicit_requirement": "明确保护要求",
    "asset_value_or_consequence": "资产价值或受损后果",
    "legal_or_standard": "法律或标准要求",
    "analyst_minimum": "分析最低保护判断",
    "inbound": "入站连接",
    "outbound": "出站连接",
    "bidirectional": "双向",
    "documented_endpoints": "文档端点",
    "connection_initiation": "连接发起关系",
    "listener_exposure": "监听/端口暴露",
    "business_data_flow": "业务数据流向",
    "to_product": "流向产品",
    "from_product": "由产品发出",
    "enabled": "已启用",
    "disabled_by_default": "默认关闭",
    "reads": "读取",
    "creates": "产生",
    "modifies": "修改",
    "transmits": "传输",
    "authenticates_with": "用于认证",
    "implements": "实现",
    "protects": "保护",
    "depends_on": "直接依赖",
    "affects": "受影响",
    "authentication": "身份认证",
    "communication": "通信",
    "cryptography": "密码技术",
    "data_collection": "数据采集",
    "spoofing": "身份欺骗",
    "tampering": "篡改",
    "repudiation": "抵赖",
    "information_disclosure": "信息泄露",
    "denial_of_service": "拒绝服务",
    "elevation_of_privilege": "权限提升",
}

_ASSET_STATUS_FILLS = {
    "confirmed": "E2F0D9",
    "conditional": "FFF2CC",
    "external_affected": "DDEBF7",
    "needs_confirmation": "FCE4D6",
}


def _value(value: Any) -> Any:
    """Convert nested Pydantic values into localized Excel-safe text."""
    if value is None:
        return ""
    if isinstance(value, Enum):
        value = value.value
    if isinstance(value, bool):
        return "是" if value else "否"
    if hasattr(value, "model_dump"):
        value = value.model_dump(mode="json")
    if isinstance(value, dict):
        rendered: list[str] = []
        for key, item in value.items():
            val_str = str(_value(item))
            if val_str == "":
                continue
            if "\n" in val_str:
                rendered.append(f"{key}:\n{val_str}")
            else:
                rendered.append(f"{key}: {val_str}")
        return "\n".join(rendered)
    if isinstance(value, (list, tuple, set)):
        items = list(value)
        if not items:
            return ""
        return "\n".join(f"- {_value(item)}" for item in items)
    if isinstance(value, str):
        return _DISPLAY_TRANSLATIONS.get(value, value)
    return value


def _tri_state_value(value: bool | None) -> str:
    """Render a three-state business answer without turning unknown into false."""
    if value is None:
        return "待确认"
    return "是" if value else "否"


def _scope_freeze_blockers(scope: Any) -> list[str]:
    """Derive review blockers from structured unknowns without asking the model to self-grade."""
    blockers: list[str] = []
    classification = scope.classification
    if scope.product_version is None:
        blockers.append("产品、硬件、固件版本及实际BOM")
    if classification.is_pde is None:
        blockers.append("CRA PDE适用性")
    if classification.classification is None:
        blockers.append("CRA产品分类（Default/Important/Critical）")
    if classification.has_rdps is None:
        blockers.append("RDPS必要性、责任边界与交付配置")
    if classification.applicable_annex is None:
        blockers.append("CRA适用附件")
    if classification.overlapping_regulations is None:
        blockers.append("RED/MD及其他欧盟法规叠加")
    if not scope.legal_basis:
        blockers.append("Node0法律依据")
    if scope.conditional_scope_components:
        blockers.append("条件组件的合同/BOM/启用状态")
    return blockers


def _pending_or_value(value: Any, pending_text: str = "待法规/业务确认") -> Any:
    return pending_text if value is None else value


def _initialize_sheet(
    sheet: Any,
    title: str,
    description: str,
    headers: list[str],
) -> None:
    """Create a consistent reader-facing Chinese sheet header."""
    sheet.append([title])
    sheet.append([description])
    sheet.append([])
    sheet.append(headers)
    if len(headers) > 1:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))


def _write_grouped_rows(
    sheet: Any,
    title: str,
    description: str,
    rows: list[tuple[str, str, Any]],
) -> None:
    """Write expanded business fields instead of nested key/value blobs."""
    _initialize_sheet(sheet, title, description, ["信息分组", "字段", "内容"])
    for group, label, value in rows:
        sheet.append([group, label, _value(value)])


def _auto_adjust_column_width(sheet: Any) -> None:
    """Adjust column widths by longest visual line using exact row/col iteration."""
    for col_idx in range(1, sheet.max_column + 1):
        max_visual = 0
        col_letter = get_column_letter(col_idx)
        for row_idx in range(4, sheet.max_row + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                text = str(cell.value)
                for line in text.split("\n"):
                    visual_len = sum(2 if ord(c) > 127 else 1 for c in line)
                    if visual_len > max_visual:
                        max_visual = visual_len

        width = max(15, min(60, max_visual + 2))
        sheet.column_dimensions[col_letter].width = width


def _style_workbook(workbook: Workbook) -> None:
    """Apply a restrained, consistent reader-facing workbook style."""
    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A5" if sheet.max_row > 4 else None
        sheet["A1"].font = Font(bold=True, size=14, color="1F1F1F")
        sheet["A2"].font = Font(italic=True, color="666666")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if sheet.max_row >= 4:
            for cell in sheet[4]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[4].height = 28
        _auto_adjust_column_width(sheet)
        # Wrapping alone does not make Excel expand rows written by openpyxl.
        for row in sheet.iter_rows(min_row=4):
            line_count = 1
            for cell in row:
                if cell.value is None:
                    continue
                width = max(8, sheet.column_dimensions[cell.column_letter].width - 3)
                lines = sum(
                    max(1, math.ceil(sum(2 if ord(c) > 127 else 1 for c in line) / width))
                    for line in str(cell.value).split("\n")
                )
                line_count = max(line_count, lines)
            sheet.row_dimensions[row[0].row].height = min(409, max(28, line_count * 16 + 8))
        sheet.row_dimensions[1].height = 25
        sheet.row_dimensions[2].height = 34


def _legal_basis_value(items: Any) -> str:
    """Render structured legal references with Chinese labels."""
    rendered: list[str] = []
    for item in items:
        values = item.model_dump(mode="json") if hasattr(item, "model_dump") else dict(item)
        parts = [f"法规/标准：{values.get('regulation', '')}"]
        for key, label in (
            ("article", "条款"),
            ("clause", "章节"),
            ("annex", "附件"),
            ("sub_clause", "子条款"),
            ("description", "说明"),
        ):
            if values.get(key):
                parts.append(f"{label}：{values[key]}")
        rendered.append("；".join(parts))
    return "\n".join(f"{index}. {text}" for index, text in enumerate(rendered, start=1))


def _write_scope(workbook: Workbook, state: dict) -> None:
    """Write Node [0] independently from Product Context."""
    scope = state["scope"]
    freeze_blockers = _scope_freeze_blockers(scope)
    sheet = workbook.active
    sheet.title = "Node0 范围定界"
    _write_grouped_rows(
        sheet,
        "Node [0] 法律与产品范围",
        "确认评估对象、CRA 属性、系统边界、依据与待确认事项。",
        [
            (
                "范围状态",
                "Node0 范围冻结状态",
                "未冻结（可继续生成资产草稿）" if freeze_blockers else "已冻结",
            ),
            ("范围状态", "正式冻结前待确认项（不阻止Node1/2草稿）", freeze_blockers),
            ("产品基础", "产品名称", scope.product_name),
            (
                "产品基础",
                "产品版本",
                _pending_or_value(scope.product_version, "待产品/BOM确认"),
            ),
            (
                "CRA 属性",
                "是否属于含数字元素产品（PDE）",
                _tri_state_value(scope.classification.is_pde),
            ),
            (
                "CRA 属性",
                "产品分类",
                _pending_or_value(scope.classification.classification),
            ),
            (
                "CRA 属性",
                "是否包含远程数据处理方案（RDPS）",
                _tri_state_value(scope.classification.has_rdps),
            ),
            ("CRA 属性", "远程数据处理方案说明", scope.classification.rdps_description),
            (
                "CRA 属性",
                "适用附件",
                _pending_or_value(scope.classification.applicable_annex),
            ),
            (
                "CRA 属性",
                "可能叠加的其他法规",
                _pending_or_value(scope.classification.overlapping_regulations),
            ),
            ("评估边界", "范围说明", scope.scope_description),
            ("评估边界", "范围内组件、接口与服务", scope.in_scope_components),
            (
                "评估边界",
                "条件适用的组件、接口与服务",
                scope.conditional_scope_components,
            ),
            ("评估边界", "范围外项目及理由", scope.out_of_scope_components),
            (
                "依据与假设",
                "法律与标准依据",
                _legal_basis_value(scope.legal_basis) or "待法规依据确认",
            ),
            ("依据与假设", "待确认问题/假设", scope.assumptions),
            ("审阅提示", "自动降级与待人工复核", scope.review_notes),
        ],
    )


def _write_product_context(workbook: Workbook, state: dict) -> None:
    """Write Node [1] summary and its functional detail sheets."""
    context = state["product_context"]
    iprfu = context.iprfu
    users = context.user_description
    environment = context.operational_environment

    summary_sheet = workbook.create_sheet("Node1 产品概况")
    _write_grouped_rows(
        summary_sheet,
        "Node [1] 产品上下文",
        "说明产品用途、用户、运行环境、已有能力和信息完整性。",
        [
            ("产品基础", "产品名称", context.product_name),
            ("产品基础", "产品版本", context.product_version),
            ("预期用途与可预见使用", "预期用途", iprfu.intended_purpose),
            ("预期用途与可预见使用", "合理可预见使用或误用", iprfu.reasonably_foreseeable_use),
            ("预期用途与可预见使用", "健康与安全考虑", iprfu.health_safety_considerations),
            ("预期用途与可预见使用", "可访问性考虑", iprfu.accessibility_considerations),
            ("预期用途与可预见使用", "支持期限", iprfu.support_period),
            ("用户", "用户类型", users.user_types),
            ("用户", "经验与能力", users.experience_level),
            ("用户", "弱势或无障碍相关群体", users.vulnerable_groups),
            ("用户", "是否作为组件集成", users.is_component),
            ("用户", "用户对远程服务的依赖", users.rdps_dependence),
            ("用户", "用户、运营方或集成商责任", users.responsibilities),
            ("运行环境", "总体说明", environment.description),
            ("运行环境", "相关网络", environment.networks),
            ("运行环境", "集成系统", environment.integrated_systems),
            ("运行环境", "物理环境", environment.physical_environment),
            ("运行环境", "网络与信任边界", environment.network_boundaries),
            ("运行环境", "运行约束", environment.constraints),
            ("运行环境", "远程服务依赖关系", environment.rdps_dependency_map),
            ("已有能力", "已有安全功能", context.existing_security_functions),
            ("已有能力", "已有非安全功能", context.existing_non_security_functions),
            ("已有能力", "远程服务依赖汇总", context.rdps_dependencies),
            ("清单完整性", "组件清单说明", context.component_inventory.completeness_notes),
            ("清单完整性", "通信矩阵说明", context.communication_matrix.completeness_notes),
            ("完整性结论", "资料完整度（不阻断资产草稿）", context.assessment),
            ("完整性结论", "结论说明", context.assessment_notes),
        ],
    )

    functions_sheet = workbook.create_sheet("Node1 产品功能")
    _initialize_sheet(
        functions_sheet,
        "Node [1] 产品功能",
        "每行是一项产品功能；安全功能用于后续资产、威胁和控制覆盖检查。",
        ["功能编号", "功能名称", "功能类别", "是否安全功能", "相关接口", "远程服务依赖", "功能说明", "限制与条件"],
    )
    for function in context.functions:
        functions_sheet.append(
            [
                _value(function.function_id),
                _value(function.name),
                _value(function.function_category),
                _value(function.is_security_function),
                _value(function.interfaces),
                _value(function.rdps_dependencies),
                _value(function.description),
                _value(function.limitations),
            ]
        )

    component_sheet = workbook.create_sheet("Node1 组件清单")
    _initialize_sheet(
        component_sheet,
        "Node [1] 组件清单",
        "每行是一项硬件、软件、固件、服务或第三方组件。",
        ["组件编号", "组件名称", "组件类型", "范围状态", "适用条件", "版本或型号", "供应方或运营方", "作用", "是否第三方", "支持期限", "相关接口"],
    )
    for component in context.component_inventory.entries:
        component_sheet.append(
            [
                _value(component.component_id),
                _value(component.name),
                _value(component.component_type),
                _value(component.scope_status),
                _value(component.conditions),
                _value(component.version),
                _value(component.supplier),
                _value(component.role),
                _value(component.is_third_party),
                _value(component.support_period),
                _value(component.interfaces),
            ]
        )

    comm_sheet = workbook.create_sheet("Node1 通信矩阵")
    _initialize_sheet(
        comm_sheet,
        "Node [1] 通信矩阵",
        "每行是一条通信流；连接方向、业务数据方向和启用状态分别记录，未知协议、端口或保护机制保持为空。",
        [
            "通信流编号",
            "源端",
            "目的端",
            "连接方向",
            "方向依据",
            "业务数据方向",
            "启用状态",
            "适用条件",
            "接口",
            "协议",
            "端口",
            "交换数据",
            "加密机制",
            "认证机制",
            "安全说明",
        ],
    )
    for entry in context.communication_matrix.entries:
        comm_sheet.append(
            [
                _value(entry.flow_id),
                _value(entry.source),
                _value(entry.destination),
                _value(entry.direction),
                _value(entry.direction_basis),
                _value(entry.business_data_direction),
                _value(entry.activation_status),
                _value(entry.conditions),
                _value(entry.interface),
                _value(entry.protocol),
                _value(entry.port),
                _value(entry.data_exchanged),
                _value(entry.encryption),
                _value(entry.authentication),
                _value(entry.security_notes),
            ]
        )


def _write_scope_and_context(workbook: Workbook, state: dict) -> None:
    """Write separate Node [0] and Node [1] reader-facing sheets."""
    _write_scope(workbook, state)
    _write_product_context(workbook, state)


def _write_assets(workbook: Workbook, state: dict) -> None:
    """Write one row per asset/objective relationship."""
    product_context = state.get("product_context")
    function_name_by_id: dict[str, str] = {}
    if product_context is not None:
        functions = (
            product_context.get("functions", [])
            if isinstance(product_context, dict)
            else product_context.functions
        )
        for function in functions:
            function_id = (
                function.get("function_id")
                if isinstance(function, dict)
                else function.function_id
            )
            function_name = (
                function.get("name")
                if isinstance(function, dict)
                else function.name
            )
            function_name_by_id[function_id] = function_name
    asset_sheet = workbook.create_sheet("Node2 资产与目标")
    _initialize_sheet(
        asset_sheet,
        "Node [2] 资产与网络安全目标",
        "每行是一组资产与保护目标；同一资产的共同字段纵向合并，目标字段仍逐行保留。",
        list(_ASSET_HEADERS),
    )
    for asset in state.get("assets", []):
        first_row = asset_sheet.max_row + 1
        for objective in asset.related_objectives:
            asset_sheet.append(
                [
                    _value(asset.asset_id),
                    _value(asset.name),
                    _value(asset.asset_type),
                    _value(asset.status),
                    _value(asset.related_function_ids),
                    _value(
                        [
                            function_name_by_id.get(function_id, function_id)
                            for function_id in asset.related_function_ids
                        ]
                    ),
                    _value(
                        [
                            (
                                f"{relationship.function_id}｜"
                                f"{_DISPLAY_TRANSLATIONS.get(relationship.relationship_type, relationship.relationship_type)}｜"
                                f"{relationship.rationale}"
                            )
                            for relationship in asset.function_relationships
                        ]
                    ),
                    _value(asset.location),
                    _value(asset.value),
                    _value(asset.description),
                    _value(objective.objective_id),
                    _value(objective.category),
                    _value(objective.basis),
                    _value(objective.status),
                    _value(objective.description),
                    _legal_basis_value(objective.legal_basis),
                ]
            )
            status_fill = _ASSET_STATUS_FILLS.get(asset.status)
            if status_fill:
                asset_sheet.cell(row=asset_sheet.max_row, column=4).fill = PatternFill(
                    "solid", fgColor=status_fill
                )
        last_row = asset_sheet.max_row
        if last_row > first_row:
            for column in range(1, 11):
                asset_sheet.merge_cells(
                    start_row=first_row,
                    start_column=column,
                    end_row=last_row,
                    end_column=column,
                )
    asset_sheet.auto_filter.ref = f"A4:P{asset_sheet.max_row}"


def _iter_evidence(value: Any, path: str = ""):
    """Yield EvidenceRef-like objects and their owning result paths recursively."""
    if value is None:
        return
    if isinstance(value, EvidenceRef):
        yield path or "$", value
        return
    if hasattr(type(value), "model_fields"):
        for field_name in value.__class__.model_fields:
            child_path = f"{path}.{field_name}" if path else field_name
            yield from _iter_evidence(getattr(value, field_name), child_path)
        return
    if isinstance(value, dict):
        for key, child in value.items():
            child_path = f"{path}.{key}" if path else str(key)
            yield from _iter_evidence(child, child_path)
        return
    if isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            yield from _iter_evidence(child, f"{path}[{index}]")


def _evidence_trace(state: dict) -> list[dict[str, Any]]:
    """Build a deduplicated, source-addressable trace across available nodes."""
    roots = {
        "scope": state.get("scope"),
        "product_context": state.get("product_context"),
        "risk_acceptance_criteria": state.get("risk_acceptance_criteria"),
        "assets": state.get("assets", []),
        "threat_assessment": state.get("threat_assessment"),
    }
    records: dict[tuple[Any, ...], dict[str, Any]] = {}
    for root_name, root_value in roots.items():
        for used_by, evidence in _iter_evidence(root_value, root_name):
            source = evidence.source
            key = (
                evidence.evidence_id,
                source.document_id,
                source.block_id,
                source.hash_sha256,
                evidence.quote,
            )
            if key not in records:
                records[key] = {
                    "evidence_id": evidence.evidence_id,
                    "used_by": [],
                    "document_id": source.document_id,
                    "block_id": source.block_id,
                    "file_name": source.file_name,
                    "file_type": source.file_type,
                    "hash_sha256": source.hash_sha256,
                    "parser_version": source.parser_version,
                    "page_number": source.page_number,
                    "printed_page_label": source.printed_page_label,
                    "section_path": list(source.section_path),
                    "quote": evidence.quote,
                    "evidence_level": _value(evidence.evidence_level),
                    "confidence": _value(evidence.confidence),
                    "retrieved_by": evidence.retrieved_by,
                    "is_human_confirmed": evidence.is_human_confirmed,
                    "notes": evidence.notes,
                }
            records[key]["used_by"].append(used_by)
    return list(records.values())


def _write_method_and_criteria(workbook: Workbook, state: dict) -> None:
    """Write the frozen method, criteria, and review conclusions near the workbook end."""
    methodology = state.get("risk_methodology")
    criteria = state.get("risk_acceptance_criteria")
    if methodology is None and criteria is None:
        return

    rows: list[tuple[str, str, Any]] = []
    if methodology is not None:
        method_labels = {
            "policy_version": "规则版本",
            "threat_modelling_method": "威胁建模方法",
            "risk_estimation_method": "风险估计方法",
            "statutory_risk_factors": "法定风险因子",
            "risk_combination_rule": "风险组合规则",
            "likelihood_dimensions": "可能性维度",
            "magnitude_dimensions": "影响程度维度",
            "damage_required_coverage": "损害评估必查范围",
            "risk_acceptance_rule": "风险接受规则",
            "treatment_priority": "风险处置优先级",
            "node_4_scale_status": "Node [4] 量表状态",
            "source_references": "方法依据",
        }
        for name, value in methodology.model_dump(mode="json").items():
            rows.append(("固定风险方法", method_labels.get(name, name), value))

    if criteria is not None:
        criteria_labels = {
            "policy_version": "规则版本",
            "regulatory_factors": "监管因素",
            "contractual_factors": "合同与供应链因素",
            "risk_nature_factors": "风险性质因素",
            "user_factors": "用户因素",
            "product_factors": "产品因素",
            "state_of_art_factors": "技术现状与社会价值因素",
            "aggregate_risk_considered": "是否考虑聚合风险",
        }
        for name, value in criteria.model_dump(mode="json", exclude={"evidence"}).items():
            rows.append(("产品特定接受准则", criteria_labels.get(name, name), value))

    if state.get("asset_assessment") is not None:
        rows.extend(
            [
                ("Node [2] 结论", "资产识别完整性", state.get("asset_assessment")),
                ("Node [2] 结论", "资产识别说明", state.get("asset_notes")),
                ("Node [2] 人工审批", "是否已批准", state.get("asset_review_approved", False)),
                ("Node [2] 人工审批", "审批备注", state.get("asset_review_notes")),
            ]
        )

    threat_assessment = state.get("threat_assessment")
    if threat_assessment is not None:
        rows.extend(
            [
                ("Node [3] 结论", "威胁覆盖完整性", threat_assessment.assessment),
                ("Node [3] 结论", "威胁建模说明", threat_assessment.notes),
            ]
        )

    sheet = workbook.create_sheet("方法与准则")
    _write_grouped_rows(
        sheet,
        "方法、接受准则与审核结论",
        "这些规则在风险分析前冻结；本表作为结果的解释与审核依据放在业务表之后。",
        rows,
    )


def _write_evidence_trace(workbook: Workbook, state: dict) -> None:
    """Write all available evidence as the final supporting sheet."""
    records = _evidence_trace(state)
    if not records:
        return
    sheet = workbook.create_sheet("证据追溯")
    _initialize_sheet(
        sheet,
        "证据追溯",
        "用于从业务结论回到客户文件、页码、章节、文本块和原文。",
        [
            "证据编号",
            "用于哪些字段",
            "文档编号",
            "文本块编号",
            "文件名",
            "PDF 物理页码",
            "文档印刷页码",
            "章节路径",
            "原文引用",
            "证据等级",
            "置信度",
            "检索方式",
            "是否人工确认",
            "备注",
        ],
    )
    for record in records:
        sheet.append(
            [
                _value(record["evidence_id"]),
                _value(record["used_by"]),
                _value(record["document_id"]),
                _value(record["block_id"]),
                _value(record["file_name"]),
                _value(record["page_number"]),
                _value(record["printed_page_label"]),
                _value(record["section_path"]),
                _value(record["quote"]),
                _value(record["evidence_level"]),
                _value(record["confidence"]),
                _value(record["retrieved_by"]),
                _value(record["is_human_confirmed"]),
                _value(record["notes"]),
            ]
        )


def _json_value(value: Any) -> Any:
    """Convert runtime values into deterministic JSON-compatible values."""
    if hasattr(value, "model_dump"):
        return value.model_dump(mode="json")
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, dict):
        return {key: _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_value(item) for item in value]
    if isinstance(value, Path):
        return str(value)
    return value


def build_progress_payload(
    state: dict,
    run_id: str,
    completed_node: str | None,
) -> dict[str, Any]:
    """Build a recoverable snapshot after one graph-node update."""
    state_keys = (
        "evidence_packet_ref",
        "scope",
        "product_context",
        "risk_methodology",
        "risk_acceptance_criteria",
        "assets",
        "asset_assessment",
        "asset_notes",
        "asset_review_approved",
        "asset_review_notes",
        "threat_assessment",
        "errors",
        "current_step",
        "status",
    )
    return {
        "schema_version": PROGRESS_SCHEMA_VERSION,
        "run_id": run_id,
        "saved_at": datetime.now(UTC).isoformat(),
        "completed_node": completed_node,
        "state": {
            key: _json_value(state.get(key))
            for key in state_keys
            if key in state
        },
        "evidence_trace": _evidence_trace(state),
    }


def _build_progress_workbook(
    state: dict,
    run_id: str,
    completed_node: str | None,
    saved_at: str,
) -> Workbook:
    """Render all currently available validated business results."""
    workbook = Workbook()
    scope = state.get("scope")
    context = state.get("product_context")

    if scope is not None and context is not None:
        _write_scope_and_context(workbook, state)
    elif scope is not None:
        _write_scope(workbook, state)
    else:
        summary_sheet = workbook.active
        summary_sheet.title = "阶段结果"
        _write_grouped_rows(
            summary_sheet,
            "工作流阶段结果",
            "当前尚无可展示的完整业务节点输出。",
            [("运行状态", "状态", state.get("status"))],
        )

    if state.get("assets"):
        _write_assets(workbook, state)

    _write_method_and_criteria(workbook, state)
    _write_evidence_trace(workbook, state)

    status_sheet = workbook.create_sheet("运行状态")
    _write_grouped_rows(
        status_sheet,
        "运行状态",
        "用于定位本次工作流执行阶段和错误，不属于正式 TARA 业务内容。",
        [
            ("运行", "运行编号", run_id),
            ("运行", "生成时间（UTC）", saved_at),
            ("运行", "最近完成节点", completed_node),
            ("运行", "当前步骤", state.get("current_step")),
            ("运行", "当前状态", state.get("status")),
            ("运行", "错误", state.get("errors", [])),
        ],
    )
    _style_workbook(workbook)
    return workbook


def export_workflow_progress(
    state: dict,
    output_dir: str | Path,
    run_id: str,
    completed_node: str | None,
    artifact_stem: str = "report",
) -> tuple[Path, Path]:
    """Atomically persist JSON and Excel after each graph-node update."""
    if not _valid_artifact_stem(artifact_stem):
        raise ValueError(f"Unsupported workflow artifact stem: {artifact_stem}")
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    payload = build_progress_payload(state, run_id, completed_node)

    status = str(state.get("status", "unknown"))
    json_path = destination / f"{artifact_stem}.json"
    excel_path = destination / f"{artifact_stem}.xlsx"
    json_temp = destination / f".{artifact_stem}.json.tmp"
    excel_temp = destination / f".{artifact_stem}.tmp.xlsx"
    json_temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    json_temp.replace(json_path)
    workbook = _build_progress_workbook(
        state,
        run_id,
        completed_node,
        payload["saved_at"],
    )
    workbook.save(excel_temp)
    excel_temp.replace(excel_path)

    logger.info("工作流进度已落盘：node=%s，status=%s，path=%s", completed_node, status, destination)
    return json_path, excel_path


def build_asset_review_payload(state: dict, run_id: str) -> dict[str, Any]:
    """Map the interrupted Node [0]-[2] state to a stable JSON review contract."""
    required = (
        "evidence_packet_ref",
        "scope",
        "product_context",
        "risk_methodology",
        "risk_acceptance_criteria",
    )
    missing = [name for name in required if state.get(name) is None]
    if missing or not state.get("assets"):
        raise ValueError(f"Cannot build asset review output; missing validated data: {missing or ['assets']}")
    if not is_asset_reviewable_state(state):
        raise ValueError(
            "Asset review output requires consistent reviewable Node [2] status and assessment"
        )
    if state.get("asset_review_approved") or state.get("threat_assessment") is not None:
        raise ValueError("Asset review output must be created before approval and Node [3]")

    def dump(item: Any) -> dict[str, Any]:
        return item.model_dump(mode="json")
    assets = list(state["assets"])
    objectives = [
        objective
        for asset in assets
        for objective in asset.related_objectives
    ]
    return {
        "schema_version": ASSET_REVIEW_SCHEMA_VERSION,
        "run_id": run_id,
        "generated_at": datetime.now(UTC).isoformat(),
        "workflow_status": state.get("status"),
        "current_step": state.get("current_step"),
        "awaiting_asset_approval": True,
        "evidence_packet_ref": dump(state["evidence_packet_ref"]),
        "scope": dump(state["scope"]),
        "product_context": dump(state["product_context"]),
        "risk_methodology": dump(state["risk_methodology"]),
        "risk_acceptance_criteria": dump(state["risk_acceptance_criteria"]),
        "asset_identification": {
            "assessment": _json_value(state.get("asset_assessment")),
            "notes": state.get("asset_notes"),
            "assets": [dump(asset) for asset in assets],
        },
        "review_summary": {
            "asset_count": len(assets),
            "objective_count": len(objectives),
            "supported_objective_count": sum(
                objective.status == "supported" for objective in objectives
            ),
            "needs_confirmation_objective_count": sum(
                objective.status == "needs_confirmation" for objective in objectives
            ),
            "typed_function_relationship_count": sum(
                len(asset.function_relationships) for asset in assets
            ),
            "conditional_asset_count": sum(
                asset.status == "conditional" for asset in assets
            ),
            "external_affected_asset_count": sum(
                asset.status == "external_affected" for asset in assets
            ),
            "needs_confirmation_asset_count": sum(
                asset.status == "needs_confirmation" for asset in assets
            ),
        },
        "review_request": {
            "review_type": "asset_confirmation",
            "step_id": "2",
            "required_action": {"approved": "boolean", "notes": "optional string"},
        },
        "evidence_trace": _evidence_trace(state),
        "errors": list(state.get("errors", [])),
    }


def _review_evidence_key(evidence: EvidenceRef) -> str:
    if evidence.source.block_id:
        return f"block:{evidence.source.block_id}"
    return f"quote:{evidence.source.file_path}:{(evidence.quote or '').strip()}"


def validate_asset_review_artifacts(
    json_path: str | Path,
    excel_path: str | Path,
) -> dict[str, int]:
    """Validate a generated Node [0]-[2] JSON/Excel pair before publication."""
    json_file = Path(json_path)
    excel_file = Path(excel_path)
    payload = json.loads(json_file.read_text(encoding="utf-8"))
    if payload.get("schema_version") != ASSET_REVIEW_SCHEMA_VERSION:
        raise ValueError("Asset review JSON uses an unsupported schema version")
    if payload.get("workflow_status") not in ASSET_REVIEWABLE_STATUSES:
        raise ValueError("Asset review JSON is not a reviewable Node [2] result")

    identification = AssetIdentificationResult.model_validate(
        payload.get("asset_identification", {})
    )
    if not is_asset_reviewable_state(
        {
            "assets": identification.assets,
            "asset_assessment": identification.assessment,
            "status": payload.get("workflow_status"),
        }
    ):
        raise ValueError(
            "Asset review JSON status does not match its asset assessment"
        )
    functions = payload.get("product_context", {}).get("functions", [])
    function_ids = {item.get("function_id") for item in functions}
    objective_ids: list[str] = []
    typed_relationship_count = 0
    for asset in identification.assets:
        unknown_ids = set(asset.related_function_ids) - function_ids
        if unknown_ids:
            raise ValueError(
                f"Asset {asset.asset_id} references unknown functions: {sorted(unknown_ids)}"
            )
        typed_ids = {
            relationship.function_id
            for relationship in asset.function_relationships
        }
        if typed_ids - function_ids:
            raise ValueError(
                f"Asset {asset.asset_id} has typed relationships to unknown functions"
            )
        if typed_ids and typed_ids != set(asset.related_function_ids):
            raise ValueError(
                f"Asset {asset.asset_id} function index does not match typed relationships"
            )
        typed_relationship_count += len(asset.function_relationships)

        asset_evidence_keys = {
            _review_evidence_key(item) for item in asset.evidence
        }
        for objective in asset.related_objectives:
            objective_ids.append(objective.objective_id)
            if not {
                _review_evidence_key(item) for item in objective.evidence
            }.issubset(asset_evidence_keys):
                raise ValueError(
                    f"Objective {objective.objective_id} cites evidence not attached to its asset"
                )
            if objective.status != "supported":
                continue
            if objective.basis in {"analyst_minimum", "unknown"}:
                raise ValueError(
                    f"Supported objective {objective.objective_id} has an unresolved basis"
                )
            if not objective.evidence:
                raise ValueError(
                    f"Supported objective {objective.objective_id} has no evidence"
                )
            if objective.basis == "asset_value_or_consequence" and not asset.value:
                raise ValueError(
                    f"Supported objective {objective.objective_id} has no asset value or consequence"
                )
            if objective.basis == "legal_or_standard" and not objective.legal_basis:
                raise ValueError(
                    f"Supported objective {objective.objective_id} has no legal or standard basis"
                )

    workbook = load_workbook(excel_file, read_only=True, data_only=True)
    try:
        missing_sheets = _ASSET_REVIEW_SHEETS - set(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(
                f"Asset review Excel is missing sheets: {sorted(missing_sheets)}"
            )
        sheet = workbook["Node2 资产与目标"]
        headers = tuple(
            sheet.cell(row=4, column=index).value
            for index in range(1, len(_ASSET_HEADERS) + 1)
        )
        if headers != _ASSET_HEADERS:
            raise ValueError("Asset review Excel Node2 headers do not match the schema")
        excel_asset_ids = [
            sheet.cell(row=row, column=1).value
            for row in range(5, sheet.max_row + 1)
            if sheet.cell(row=row, column=1).value
        ]
        excel_objective_ids = [
            sheet.cell(row=row, column=11).value
            for row in range(5, sheet.max_row + 1)
            if sheet.cell(row=row, column=11).value
        ]
    finally:
        workbook.close()

    asset_ids = [asset.asset_id for asset in identification.assets]
    if excel_asset_ids != asset_ids:
        raise ValueError("Asset review JSON and Excel asset identifiers differ")
    if excel_objective_ids != objective_ids:
        raise ValueError("Asset review JSON and Excel objective identifiers differ")

    summary = payload.get("review_summary", {})
    expected_summary = {
        "asset_count": len(asset_ids),
        "objective_count": len(objective_ids),
        "supported_objective_count": sum(
            objective.status == "supported"
            for asset in identification.assets
            for objective in asset.related_objectives
        ),
        "needs_confirmation_objective_count": sum(
            objective.status == "needs_confirmation"
            for asset in identification.assets
            for objective in asset.related_objectives
        ),
        "typed_function_relationship_count": typed_relationship_count,
        "conditional_asset_count": sum(
            asset.status == "conditional" for asset in identification.assets
        ),
        "external_affected_asset_count": sum(
            asset.status == "external_affected" for asset in identification.assets
        ),
        "needs_confirmation_asset_count": sum(
            asset.status == "needs_confirmation" for asset in identification.assets
        ),
    }
    for key, value in expected_summary.items():
        if summary.get(key) != value:
            raise ValueError(f"Asset review summary is inconsistent: {key}")
    return expected_summary


def export_asset_review(
    state: dict,
    output_dir: str | Path,
    run_id: str,
    artifact_stem: str = "report",
) -> tuple[Path, Path]:
    """Export the pre-approval Node [0]-[2] review package as JSON and Excel."""
    if not _valid_artifact_stem(artifact_stem, asset_review=True):
        raise ValueError(f"Unsupported asset review artifact stem: {artifact_stem}")
    payload = build_asset_review_payload(state, run_id)
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    json_path = destination / f"{artifact_stem}.json"
    excel_path = destination / f"{artifact_stem}.xlsx"

    json_temp = destination / f".{artifact_stem}.json.tmp"
    json_temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    excel_temp = destination / f".{artifact_stem}.tmp.xlsx"
    try:
        workbook = Workbook()
        _write_scope_and_context(workbook, state)
        _write_assets(workbook, state)
        _write_method_and_criteria(workbook, state)
        _write_evidence_trace(workbook, state)
        _style_workbook(workbook)
        workbook.save(excel_temp)
        audit = validate_asset_review_artifacts(json_temp, excel_temp)
        json_temp.replace(json_path)
        excel_temp.replace(excel_path)
    except Exception:
        json_temp.unlink(missing_ok=True)
        excel_temp.unlink(missing_ok=True)
        raise
    logger.info(
        "Node [0]-[2] 审核包已导出并通过一致性检查: %s (assets=%d, objectives=%d, typed_relationships=%d)",
        destination,
        audit["asset_count"],
        audit["objective_count"],
        audit["typed_function_relationship_count"],
    )
    return json_path, excel_path


def export_tara_report(state: dict, output_path: str = "tara_report.xlsx") -> None:
    """Serialize validated terminal state into the current multi-sheet report."""
    scope = state.get("scope")
    product_context = state.get("product_context")
    if scope is None or product_context is None:
        logger.error("导出失败：State 缺少 scope 或 product_context")
        return

    workbook = Workbook()
    _write_scope_and_context(workbook, state)

    assets = state.get("assets", [])
    if assets:
        _write_assets(workbook, state)

    threat_assessment = state.get("threat_assessment")
    if threat_assessment is not None and threat_assessment.threats:
        threat_sheet = workbook.create_sheet("Node3 STRIDE 威胁")
        threat_headers = [
            "威胁编号",
            "STRIDE 类别",
            "威胁名称",
            "目标资产编号",
            "受损安全目标编号",
            "受损原因",
            "威胁来源或攻击者",
            "攻击路径或前提",
            "已知可利用漏洞",
            "威胁场景说明",
            "置信度",
        ]
        _initialize_sheet(
            threat_sheet,
            "Node [3] STRIDE 威胁场景",
            "只针对人工批准的资产和网络安全目标生成威胁；DREAD 评分属于 Node [4]。",
            threat_headers,
        )

        for threat in threat_assessment.threats:
            threat_sheet.append(
                [
                    _value(threat.threat_id),
                    _value(threat.stride_category),
                    _value(threat.title),
                    _value([asset.asset_id for asset in threat.targeted_assets]),
                    _value(
                        [objective.objective_id for objective in threat.compromised_objectives]
                    ),
                    _value(threat.cause_of_compromise),
                    _value(threat.threat_actor),
                    _value(threat.attack_path),
                    _value(threat.known_vulnerabilities),
                    _value(threat.description),
                    _value(threat.confidence),
                ]
            )

    _write_method_and_criteria(workbook, state)
    _write_evidence_trace(workbook, state)
    _style_workbook(workbook)

    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    logger.info("TARA Excel 报告已导出: %s", destination)
